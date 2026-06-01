"""
Extract and score text from candidate documents.

This layer is deliberately lazy: it only opens files after metadata search has
selected candidates. Extracted text is cached under indexes/ so repeated local
experiments do not keep re-parsing the same PDFs.
"""

from __future__ import annotations

import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from document_store import resolve_document_path
from search import SearchResult, has_cjk, normalize_text, tokenize


DEFAULT_CONTENT_CACHE_PATH = Path("indexes/content_cache.json")
EXTRACTOR_VERSION = "office-ocr-v1"
DOCX_CHUNK_CHARS = 2500
MAX_EXCEL_ROWS_PER_SHEET = 200
MAX_EXCEL_COLS_PER_SHEET = 50
OCR_MIN_CHARS_PER_FILE = 120
OCR_LANGUAGES = "eng"
OCR_RENDER_DPI = 160
OCR_TIMEOUT_SECONDS_PER_PAGE = 30
SUPPORTED_CONTENT_EXTENSIONS = {".pdf", ".pptx", ".docx", ".xlsx", ".xlsm"}

CONTENT_SCORE_STOPWORDS = {
    "about",
    "deck",
    "doc",
    "document",
    "file",
    "mentions",
    "pdf",
    "presentation",
    "probably",
    "report",
    "says",
    "slide",
    "slides",
    "talk",
    "talks",
    "topic",
    "workbook",
    "worksheet",
}


@dataclass
class ContentEvidence:
    file_id: str
    score: float
    matched_terms: list[str]
    snippets: list[str]
    inspected_pages: int
    page_count: int
    cache_hit: bool
    error: str | None = None
    ocr_used: bool = False
    ocr_pages: int = 0
    ocr_error: str | None = None
    text_samples: list[str] = field(default_factory=list)


def inspect_content_for_results(
    query: str,
    results: list[SearchResult],
    max_files: int,
    max_pages_per_file: int,
    max_chars_per_file: int,
    cache_path: Path = DEFAULT_CONTENT_CACHE_PATH,
) -> list[ContentEvidence]:
    cache = load_cache(cache_path)
    evidence: list[ContentEvidence] = []

    for result in results[:max_files]:
        record = result.record
        extension = str(record.get("extension", "")).lower()
        if extension not in SUPPORTED_CONTENT_EXTENSIONS:
            evidence.append(
                ContentEvidence(
                    file_id=str(record.get("file_id")),
                    score=0,
                    matched_terms=[],
                    snippets=[],
                    inspected_pages=0,
                    page_count=0,
                    cache_hit=False,
                    error=f"content extraction for {extension} is not supported yet",
                )
            )
            continue

        try:
            extracted = get_or_extract_document_text(
                record=record,
                cache=cache,
                max_units=max_pages_per_file,
                max_chars=max_chars_per_file,
            )
        except Exception as error:
            evidence.append(
                ContentEvidence(
                    file_id=str(record.get("file_id")),
                    score=0,
                    matched_terms=[],
                    snippets=[],
                    inspected_pages=0,
                    page_count=0,
                    cache_hit=False,
                    error=format_content_extraction_error(record, error),
                )
            )
            continue

        evidence.append(score_extracted_text(query, extracted))

    save_cache(cache_path, cache)
    return evidence


def get_or_extract_document_text(
    record: dict[str, Any],
    cache: dict[str, Any],
    max_units: int,
    max_chars: int,
) -> dict[str, Any]:
    file_id = str(record.get("file_id"))
    cached = cache.get(file_id)
    cache_key = {
        "extractor_version": EXTRACTOR_VERSION,
        "source": record.get("source"),
        "uri": record.get("uri") or record.get("relative_path"),
        "extension": record.get("extension"),
        "file_size_bytes": record.get("file_size_bytes"),
        "modified_time": record.get("modified_time"),
        "max_units": max_units,
        "max_chars": max_chars,
        "ocr_min_chars_per_file": get_ocr_min_chars_per_file(),
        "ocr_languages": get_ocr_languages(),
    }

    if cached and cached.get("cache_key") == cache_key:
        cached["cache_hit"] = True
        return cached

    extracted = extract_document_text(
        record,
        max_units=max_units,
        max_chars=max_chars,
    )
    extracted["cache_key"] = cache_key
    extracted["cache_hit"] = False
    cache[file_id] = extracted
    return extracted



def format_content_extraction_error(record: dict[str, Any], error: Exception) -> str:
    filename = str(record.get("filename") or record.get("uri") or "candidate file")
    extension = str(record.get("extension", "")).lower()

    if isinstance(error, zipfile.BadZipFile) or "not a zip file" in str(error).lower():
        return (
            f"Could not read {filename} as {extension or 'an Office file'} because "
            "it is not a valid zip-based Office document. Legacy .doc/.ppt/.xls "
            "or mislabeled files should be converted to .docx/.pptx/.xlsx or PDF."
        )

    return f"Could not extract text from {filename}: {error}"


def extract_document_text(
    record: dict[str, Any],
    max_units: int,
    max_chars: int,
) -> dict[str, Any]:
    extension = str(record.get("extension", "")).lower()

    if extension == ".pdf":
        extracted = extract_pdf_text(record, max_pages=max_units, max_chars=max_chars)
    elif extension == ".pptx":
        extracted = extract_pptx_text(record, max_slides=max_units, max_chars=max_chars)
    elif extension == ".docx":
        extracted = extract_docx_text(record, max_chunks=max_units, max_chars=max_chars)
    elif extension in {".xlsx", ".xlsm"}:
        extracted = extract_xlsx_text(record, max_sheets=max_units, max_chars=max_chars)
    else:
        raise ValueError(f"content extraction for {extension} is not supported yet")

    return add_local_ocr_fallback_if_needed(
        record=record,
        extracted=extracted,
        max_units=max_units,
        max_chars=max_chars,
    )


def extract_pdf_text(
    record: dict[str, Any],
    max_pages: int,
    max_chars: int,
) -> dict[str, Any]:
    try:
        import fitz
    except ImportError as error:
        raise RuntimeError(
            "Content inspection needs PyMuPDF. Run: .venv/bin/python -m pip install -r requirements.txt"
        ) from error

    path = resolve_document_path(record)
    pages: list[dict[str, Any]] = []
    total_chars = 0

    with fitz.open(path) as document:
        page_count = document.page_count
        for page_index in range(min(page_count, max_pages)):
            if total_chars >= max_chars:
                break

            page = document.load_page(page_index)
            text = page.get_text("text")
            remaining_chars = max_chars - total_chars
            text = text[:remaining_chars]
            total_chars += len(text)
            pages.append(
                {
                    "page_number": page_index + 1,
                    "text": text,
                }
            )

    return {
        "file_id": record.get("file_id"),
        "filename": record.get("filename"),
        "unit_label": "page",
        "page_count": page_count,
        "inspected_pages": len(pages),
        "char_count": total_chars,
        "pages": pages,
    }


def extract_pptx_text(
    record: dict[str, Any],
    max_slides: int,
    max_chars: int,
) -> dict[str, Any]:
    try:
        from pptx import Presentation
    except ImportError as error:
        raise RuntimeError(
            "PPTX content inspection needs python-pptx. Run: .venv/bin/python -m pip install -r requirements.txt"
        ) from error

    path = resolve_document_path(record)
    presentation = Presentation(path)
    pages: list[dict[str, Any]] = []
    total_chars = 0

    for slide_index, slide in enumerate(presentation.slides, start=1):
        if len(pages) >= max_slides or total_chars >= max_chars:
            break

        text = "\n".join(extract_slide_text(slide))
        remaining_chars = max_chars - total_chars
        text = text[:remaining_chars]
        total_chars += len(text)
        pages.append(
            {
                "page_number": slide_index,
                "text": text,
            }
        )

    return {
        "file_id": record.get("file_id"),
        "filename": record.get("filename"),
        "unit_label": "slide",
        "page_count": len(presentation.slides),
        "inspected_pages": len(pages),
        "char_count": total_chars,
        "pages": pages,
    }


def extract_slide_text(slide: Any) -> list[str]:
    parts: list[str] = []

    for shape in slide.shapes:
        text = getattr(shape, "text", "")
        if text:
            parts.append(text)

        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                row_values = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text and cell.text.strip()
                ]
                if row_values:
                    parts.append(" | ".join(row_values))

    if getattr(slide, "has_notes_slide", False):
        notes_text_frame = getattr(slide.notes_slide, "notes_text_frame", None)
        if notes_text_frame and notes_text_frame.text:
            parts.append("Speaker notes: " + notes_text_frame.text)

    return parts


def extract_docx_text(
    record: dict[str, Any],
    max_chunks: int,
    max_chars: int,
) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError as error:
        raise RuntimeError(
            "DOCX content inspection needs python-docx. Run: .venv/bin/python -m pip install -r requirements.txt"
        ) from error

    path = resolve_document_path(record)
    document = Document(path)
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            row_values = [
                cell.text.strip()
                for cell in row.cells
                if cell.text and cell.text.strip()
            ]
            if row_values:
                parts.append(" | ".join(row_values))

    full_text = "\n".join(parts)[:max_chars]
    chunks = split_text_into_units(full_text, unit_size=DOCX_CHUNK_CHARS)[:max_chunks]
    pages = [
        {
            "page_number": index,
            "text": chunk,
        }
        for index, chunk in enumerate(chunks, start=1)
    ]

    return {
        "file_id": record.get("file_id"),
        "filename": record.get("filename"),
        "unit_label": "chunk",
        "page_count": len(chunks),
        "inspected_pages": len(pages),
        "char_count": sum(len(page["text"]) for page in pages),
        "pages": pages,
    }


def extract_xlsx_text(
    record: dict[str, Any],
    max_sheets: int,
    max_chars: int,
) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "XLSX content inspection needs openpyxl. Run: .venv/bin/python -m pip install -r requirements.txt"
        ) from error

    path = resolve_document_path(record)
    workbook = load_workbook(path, read_only=True, data_only=False)
    pages: list[dict[str, Any]] = []
    total_chars = 0

    for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
        if len(pages) >= max_sheets or total_chars >= max_chars:
            break

        lines = [f"Sheet: {worksheet.title}"]
        for row in worksheet.iter_rows(
            max_row=min(worksheet.max_row or 0, MAX_EXCEL_ROWS_PER_SHEET),
            max_col=min(worksheet.max_column or 0, MAX_EXCEL_COLS_PER_SHEET),
        ):
            cell_values = []
            for cell in row:
                if cell.value is None:
                    continue
                cell_values.append(f"{cell.coordinate}={cell.value}")
            if cell_values:
                lines.append(" | ".join(cell_values))

            if len("\n".join(lines)) >= max_chars - total_chars:
                break

        text = "\n".join(lines)
        remaining_chars = max_chars - total_chars
        text = text[:remaining_chars]
        total_chars += len(text)
        pages.append(
            {
                "page_number": sheet_index,
                "text": text,
            }
        )

    workbook.close()

    return {
        "file_id": record.get("file_id"),
        "filename": record.get("filename"),
        "unit_label": "sheet",
        "page_count": len(workbook.worksheets),
        "inspected_pages": len(pages),
        "char_count": total_chars,
        "pages": pages,
    }


def add_local_ocr_fallback_if_needed(
    record: dict[str, Any],
    extracted: dict[str, Any],
    max_units: int,
    max_chars: int,
) -> dict[str, Any]:
    char_count = int(extracted.get("char_count", 0))
    minimum_chars = get_ocr_min_chars_per_file()
    if char_count >= minimum_chars:
        extracted["ocr"] = {
            "attempted": False,
            "used": False,
            "reason": f"extracted text already had {char_count} characters",
        }
        return extracted

    extension = str(record.get("extension", "")).lower()
    if extension != ".pdf":
        extracted["ocr"] = {
            "attempted": False,
            "used": False,
            "reason": (
                f"local OCR fallback for {extension} needs a page renderer; "
                "the current MVP supports scanned PDF pages"
            ),
        }
        return extracted

    remaining_chars = max(max_chars - char_count, 0)
    if remaining_chars <= 0:
        extracted["ocr"] = {
            "attempted": False,
            "used": False,
            "reason": "character budget was already exhausted",
        }
        return extracted

    try:
        ocr_result = ocr_pdf_pages(
            record=record,
            max_pages=max_units,
            max_chars=remaining_chars,
        )
    except RuntimeError as error:
        extracted["ocr"] = {
            "attempted": True,
            "used": False,
            "error": str(error),
        }
        return extracted

    pages_by_number = {
        int(page.get("page_number", 0)): page
        for page in extracted.get("pages", [])
    }
    ocr_pages = ocr_result["pages"]
    for page in ocr_pages:
        page_number = int(page["page_number"])
        ocr_text = str(page["text"]).strip()
        if not ocr_text:
            continue

        existing_page = pages_by_number.get(page_number)
        if existing_page is None:
            existing_page = {
                "page_number": page_number,
                "text": "",
            }
            extracted.setdefault("pages", []).append(existing_page)
            pages_by_number[page_number] = existing_page

        existing_text = str(existing_page.get("text", ""))
        separator = "\n\n" if existing_text else ""
        existing_page["text"] = f"{existing_text}{separator}[OCR text]\n{ocr_text}"

    used = bool(ocr_pages)
    extracted["ocr"] = {
        "attempted": True,
        "used": used,
        "pages": [page["page_number"] for page in ocr_pages],
        "source": "local_tesseract",
        "languages": get_ocr_languages(),
        "errors": ocr_result["errors"],
    }
    extracted["char_count"] = sum(
        len(str(page.get("text", "")))
        for page in extracted.get("pages", [])
    )
    extracted["inspected_pages"] = len(extracted.get("pages", []))
    return extracted


def ocr_pdf_pages(
    record: dict[str, Any],
    max_pages: int,
    max_chars: int,
) -> dict[str, Any]:
    tesseract_path = shutil.which("tesseract")
    if not tesseract_path:
        raise RuntimeError(
            "local OCR fallback needs the Tesseract command-line tool. "
            "Install it locally, or keep relying on normal text extraction."
        )

    try:
        import fitz
    except ImportError as error:
        raise RuntimeError(
            "PDF OCR fallback needs PyMuPDF. Run: .venv/bin/python -m pip install -r requirements.txt"
        ) from error

    path = resolve_document_path(record)
    languages = get_ocr_languages()
    pages: list[dict[str, Any]] = []
    errors: list[str] = []
    total_chars = 0

    with fitz.open(path) as document, tempfile.TemporaryDirectory() as tmp_dir:
        page_count = min(document.page_count, max_pages)
        zoom = OCR_RENDER_DPI / 72
        matrix = fitz.Matrix(zoom, zoom)

        for page_index in range(page_count):
            if total_chars >= max_chars:
                break

            page = document.load_page(page_index)
            image_path = Path(tmp_dir) / f"page-{page_index + 1}.png"
            pixmap = page.get_pixmap(matrix=matrix, alpha=False)
            pixmap.save(image_path)

            try:
                completed = subprocess.run(
                    [
                        tesseract_path,
                        str(image_path),
                        "stdout",
                        "-l",
                        languages,
                    ],
                    capture_output=True,
                    text=True,
                    timeout=OCR_TIMEOUT_SECONDS_PER_PAGE,
                    check=False,
                )
            except subprocess.TimeoutExpired:
                errors.append(f"page {page_index + 1}: OCR timed out")
                continue
            if completed.returncode != 0:
                message = completed.stderr.strip() or "unknown OCR error"
                errors.append(f"page {page_index + 1}: {message}")
                continue

            text = completed.stdout.strip()
            if not text:
                continue

            remaining_chars = max_chars - total_chars
            text = text[:remaining_chars]
            total_chars += len(text)
            pages.append(
                {
                    "page_number": page_index + 1,
                    "text": text,
                }
            )

    return {
        "pages": pages,
        "errors": errors,
    }


def get_ocr_min_chars_per_file() -> int:
    return OCR_MIN_CHARS_PER_FILE


def get_ocr_languages() -> str:
    return OCR_LANGUAGES


def split_text_into_units(text: str, unit_size: int) -> list[str]:
    if not text:
        return []

    return [
        text[index : index + unit_size]
        for index in range(0, len(text), unit_size)
    ]


def score_extracted_text(query: str, extracted: dict[str, Any]) -> ContentEvidence:
    file_id = str(extracted.get("file_id"))
    query_terms = [
        term
        for term in tokenize(query)
        if (len(term) >= 3 or has_cjk(term)) and term not in CONTENT_SCORE_STOPWORDS
    ]
    matched_terms: set[str] = set()
    snippets: list[str] = []

    for page in extracted.get("pages", []):
        page_number = int(page.get("page_number", 0))
        text = str(page.get("text", ""))
        normalized_text = normalize_for_content(text)

        for term in query_terms:
            if not term_matches(normalized_text, term):
                continue

            matched_terms.add(term)
            unit_label = str(extracted.get("unit_label", "page"))
            snippet = make_snippet(text, term)
            if snippet:
                snippets.append(f"{unit_label} {page_number}: {snippet}")

    unique_snippets = dedupe(snippets)[:3]
    score = float(len(matched_terms) * 6 + len(unique_snippets) * 2)
    ocr = extracted.get("ocr") or {}
    ocr_pages = ocr.get("pages") or []
    ocr_error = ocr.get("error")
    ocr_errors = ocr.get("errors") or []
    if not ocr_error and ocr_errors:
        ocr_error = "; ".join(str(error) for error in ocr_errors[:2])

    return ContentEvidence(
        file_id=file_id,
        score=score,
        matched_terms=sorted(matched_terms),
        snippets=unique_snippets,
        text_samples=make_text_samples(extracted, query_terms),
        inspected_pages=int(extracted.get("inspected_pages", 0)),
        page_count=int(extracted.get("page_count", 0)),
        cache_hit=bool(extracted.get("cache_hit", False)),
        ocr_used=bool(ocr.get("used", False)),
        ocr_pages=len(ocr_pages),
        ocr_error=ocr_error,
    )


def normalize_for_content(text: str) -> str:
    return normalize_text(text)


def term_matches(normalized_text: str, term: str) -> bool:
    if has_cjk(term):
        return term in normalized_text

    pattern = rf"\b{re.escape(term)}[a-z0-9]*\b"
    return re.search(pattern, normalized_text) is not None


def make_text_samples(
    extracted: dict[str, Any],
    query_terms: list[str],
    max_samples: int = 4,
    max_chars: int = 700,
) -> list[str]:
    samples: list[str] = []
    pages = extracted.get("pages", [])
    unit_label = str(extracted.get("unit_label", "page"))

    matching_pages = [
        page
        for page in pages
        if any(
            term_matches(normalize_for_content(str(page.get("text", ""))), term)
            for term in query_terms
        )
    ]
    sample_pages = matching_pages or pages

    for page in sample_pages:
        text = compact_text(str(page.get("text", "")))
        if not text:
            continue
        page_number = int(page.get("page_number", 0))
        samples.append(f"{unit_label} {page_number}: {text[:max_chars]}")
        if len(samples) >= max_samples:
            break

    return samples


def compact_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def make_snippet(text: str, term: str, window: int = 80) -> str | None:
    if has_cjk(term):
        match = re.search(re.escape(term), text, re.IGNORECASE)
        if not match:
            normalized = normalize_for_content(text)
            match = re.search(re.escape(term), normalized, re.IGNORECASE)
            if not match:
                return None
            text = normalized
        return make_snippet_from_match(text, match, window)

    match = re.search(rf"\b{re.escape(term)}[a-z0-9]*\b", text, re.IGNORECASE)
    if not match:
        return None

    return make_snippet_from_match(text, match, window)


def make_snippet_from_match(text: str, match: re.Match[str], window: int) -> str:
    start = max(match.start() - window, 0)
    end = min(match.end() + window, len(text))
    snippet = re.sub(r"\s+", " ", text[start:end]).strip()
    if start > 0:
        snippet = "..." + snippet
    if end < len(text):
        snippet = snippet + "..."
    return snippet


def load_cache(cache_path: Path) -> dict[str, Any]:
    if not cache_path.exists():
        return {}

    with cache_path.open("r", encoding="utf-8") as f:
        loaded = json.load(f)

    if isinstance(loaded, dict):
        return loaded
    return {}


def save_cache(cache_path: Path, cache: dict[str, Any]) -> None:
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    with cache_path.open("w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)


def dedupe(values: list[str]) -> list[str]:
    seen: set[str] = set()
    output: list[str] = []

    for value in values:
        if value in seen:
            continue
        seen.add(value)
        output.append(value)

    return output
